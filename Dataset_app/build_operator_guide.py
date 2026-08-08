from __future__ import annotations

import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\CV_daech")
OUT_DOCX = ROOT / "operator_guide.docx"
ASSET_DIR = ROOT / "_operator_guide_assets"
WORKFLOW_IMG = ASSET_DIR / "workflow.png"
UI_MAP_IMG = ASSET_DIR / "ui_map.png"


def find_font(names: list[str], size: int) -> ImageFont.FreeTypeFont:
    candidates = []
    for name in names:
        if Path(name).is_file():
            candidates.append(Path(name))

    if not candidates:
        font_dirs = [
            Path(r"C:\Windows\Fonts"),
            Path(r"C:\Users\kaykov\.fonts"),
        ]
        for font_dir in font_dirs:
            if not font_dir.exists():
                continue
            for name in names:
                candidate = font_dir / name
                if candidate.exists():
                    candidates.append(candidate)
                    break
            if candidates:
                break

    if not candidates:
        raise FileNotFoundError(f"Unable to locate a font from {names}")

    return ImageFont.truetype(str(candidates[0]), size=size)


FONT_REG = find_font(["arial.ttf", "DejaVuSans.ttf"], 24)
FONT_BOLD = find_font(["arialbd.ttf", "DejaVuSans-Bold.ttf"], 24)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        test = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), test, font=font)[2] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines or [text]


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color: str, width: int = 6) -> None:
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    head_len = 16
    head_w = 10
    p1 = end
    p2 = (
        int(end[0] - head_len * math.cos(angle) + head_w * math.sin(angle)),
        int(end[1] - head_len * math.sin(angle) - head_w * math.cos(angle)),
    )
    p3 = (
        int(end[0] - head_len * math.cos(angle) - head_w * math.sin(angle)),
        int(end[1] - head_len * math.sin(angle) + head_w * math.cos(angle)),
    )
    draw.polygon([p1, p2, p3], fill=color)


def rounded_box(draw: ImageDraw.ImageDraw, box, fill, outline, radius=24, width=3):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def make_workflow_image(path: Path) -> None:
    w, h = 1800, 720
    img = Image.new("RGB", (w, h), "#F7F9FC")
    draw = ImageDraw.Draw(img)

    title_font = find_font(["arialbd.ttf", "DejaVuSans-Bold.ttf"], 42)
    subtitle_font = find_font(["arial.ttf", "DejaVuSans.ttf"], 22)
    step_font = find_font(["arialbd.ttf", "DejaVuSans-Bold.ttf"], 25)
    body_font = find_font(["arial.ttf", "DejaVuSans.ttf"], 19)

    draw.text((70, 45), "Быстрый маршрут оператора", font=title_font, fill="#143A66")
    draw.text((70, 100), "Пять основных действий от выбора плана до сохранения снимка", font=subtitle_font, fill="#58667A")

    steps = [
        ("1", "Выбрать план", "Указать JSON-конфигурацию"),
        ("2", "Выбрать датасет", "Папка сохранения результата"),
        ("3", "Загрузить план", "Программа построит структуру"),
        ("4", "Подключить камеру", "Daheng или тестовый режим"),
        ("5", "Сделать снимок", "Сохранение кадра и метаданных"),
    ]

    colors = ["#DCEBFA", "#EAF4E1", "#FFF1D9", "#EFE3FF", "#FBE3E3"]
    outlines = ["#5A92C7", "#7BA55A", "#D1A44D", "#9368C8", "#C56E6E"]
    xs = [70, 395, 720, 1045, 1370]
    box_y = 245
    box_w = 290
    box_h = 220

    for idx, ((num, header, desc), fill, outline, x) in enumerate(zip(steps, colors, outlines, xs)):
        rounded_box(draw, (x, box_y, x + box_w, box_y + box_h), fill, outline)
        circ_center = (x + 44, box_y + 42)
        draw.ellipse((circ_center[0] - 26, circ_center[1] - 26, circ_center[0] + 26, circ_center[1] + 26), fill=outline)
        nfont = find_font(["arialbd.ttf", "DejaVuSans-Bold.ttf"], 24)
        nbbox = draw.textbbox((0, 0), num, font=nfont)
        draw.text((circ_center[0] - (nbbox[2] - nbbox[0]) / 2, circ_center[1] - 15), num, font=nfont, fill="white")
        draw.text((x + 82, box_y + 24), header, font=step_font, fill="#17324D")

        lines = wrap_text(draw, desc, body_font, box_w - 42)
        ty = box_y + 85
        for line in lines:
            draw.text((x + 26, ty), line, font=body_font, fill="#374151")
            ty += 28

        if idx < len(steps) - 1:
            start = (x + box_w + 10, box_y + box_h // 2)
            end = (xs[idx + 1] - 18, box_y + box_h // 2)
            draw_arrow(draw, start, end, "#55789E", width=7)

    draw.rounded_rectangle((70, 520, 1730, 665), radius=24, fill="white", outline="#D9E2EC", width=2)
    note_font = find_font(["arialbd.ttf", "DejaVuSans-Bold.ttf"], 24)
    draw.text((95, 545), "Подсказка", font=note_font, fill="#1F4D78")
    note_text = "Если камера Daheng не доступна, переключитесь на режим «Тестовая» и проверьте путь к папке датасета."
    note_lines = wrap_text(draw, note_text, body_font, 1550)
    y = 585
    for line in note_lines:
        draw.text((95, y), line, font=body_font, fill="#334155")
        y += 28

    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def make_ui_map_image(path: Path) -> None:
    w, h = 1800, 960
    img = Image.new("RGB", (w, h), "#F7F9FC")
    draw = ImageDraw.Draw(img)

    title_font = find_font(["arialbd.ttf", "DejaVuSans-Bold.ttf"], 42)
    subtitle_font = find_font(["arial.ttf", "DejaVuSans.ttf"], 22)
    label_font = find_font(["arialbd.ttf", "DejaVuSans-Bold.ttf"], 22)
    body_font = find_font(["arial.ttf", "DejaVuSans.ttf"], 19)

    draw.text((70, 45), "Окно программы: что где находится", font=title_font, fill="#143A66")
    draw.text((70, 100), "Схема помогает быстро сориентироваться перед началом съёмки", font=subtitle_font, fill="#58667A")

    # Main app shell
    app_x, app_y = 140, 220
    app_w, app_h = 1320, 630
    draw.rounded_rectangle((app_x, app_y, app_x + app_w, app_y + app_h), radius=30, fill="white", outline="#CBD5E1", width=3)
    draw.rectangle((app_x, app_y, app_x + app_w, app_y + 54), fill="#E9F0F8")
    draw.text((app_x + 24, app_y + 14), "CassetteDatasetCapture", font=label_font, fill="#17324D")

    # Sidebar
    sb_x, sb_y, sb_w, sb_h = app_x + 24, app_y + 72, 320, 500
    rounded_box(draw, (sb_x, sb_y, sb_x + sb_w, sb_y + sb_h), "#F8FBFF", "#D7E3F1", radius=18, width=2)
    draw.text((sb_x + 18, sb_y + 18), "Панель управления", font=label_font, fill="#26415D")
    draw.rounded_rectangle((sb_x + 18, sb_y + 62, sb_x + 300, sb_y + 104), radius=10, fill="#EEF4FA", outline="#C3D4E4", width=2)
    draw.text((sb_x + 32, sb_y + 72), "Путь к конфигурации", font=body_font, fill="#314155")
    draw.rounded_rectangle((sb_x + 18, sb_y + 125, sb_x + 300, sb_y + 167), radius=10, fill="#EEF4FA", outline="#C3D4E4", width=2)
    draw.text((sb_x + 32, sb_y + 135), "Папка датасета", font=body_font, fill="#314155")
    draw.rounded_rectangle((sb_x + 18, sb_y + 190, sb_x + 300, sb_y + 232), radius=10, fill="#EEF4FA", outline="#C3D4E4", width=2)
    draw.text((sb_x + 32, sb_y + 200), "Загрузить план", font=body_font, fill="#314155")
    draw.rounded_rectangle((sb_x + 18, sb_y + 250, sb_x + 300, sb_y + 292), radius=10, fill="#EEF4FA", outline="#C3D4E4", width=2)
    draw.text((sb_x + 32, sb_y + 260), "Подключить камеру", font=body_font, fill="#314155")
    draw.rounded_rectangle((sb_x + 18, sb_y + 310, sb_x + 300, sb_y + 352), radius=10, fill="#EEF4FA", outline="#C3D4E4", width=2)
    draw.text((sb_x + 32, sb_y + 320), "Сделать снимок", font=body_font, fill="#314155")

    # Preview area
    pr_x, pr_y, pr_w, pr_h = app_x + 370, app_y + 72, 870, 290
    rounded_box(draw, (pr_x, pr_y, pr_x + pr_w, pr_y + pr_h), "#101317", "#334155", radius=18, width=2)
    draw.text((pr_x + 24, pr_y + 18), "Окно предпросмотра камеры", font=label_font, fill="#DDE7F3")
    draw.rounded_rectangle((pr_x + 20, pr_y + 60, pr_x + 850, pr_y + 250), radius=14, fill="#1D232C", outline="#3A4454", width=2)
    draw.text((pr_x + 50, pr_y + 132), "Здесь видно текущий кадр", font=body_font, fill="#A8B3C2")

    # Slot map / status zone
    grid_x, grid_y = app_x + 370, app_y + 388
    rounded_box(draw, (grid_x, grid_y, grid_x + 870, grid_y + 214), "#F8FBFF", "#D7E3F1", radius=18, width=2)
    draw.text((grid_x + 24, grid_y + 16), "Схема слотов и статус текущей конфигурации", font=label_font, fill="#26415D")
    for row in range(2):
        for col in range(5):
            x = grid_x + 24 + col * 160
            y = grid_y + 58 + row * 68
            draw.rounded_rectangle((x, y, x + 134, y + 50), radius=10, fill="#EAF4E1" if (row + col) % 2 == 0 else "#FFF1D9", outline="#B8C6D6", width=2)
            draw.text((x + 50, y + 14), f"{row * 5 + col + 1}", font=body_font, fill="#2F3C4D")

    # Callout boxes and arrows
    callouts = [
        ((30, 260, 120, 320), "1", "#5A92C7"),
        ((30, 345, 120, 405), "2", "#7BA55A"),
        ((1530, 320, 1720, 380), "3", "#D1A44D"),
        ((1530, 500, 1730, 560), "4", "#9368C8"),
        ((1530, 690, 1720, 750), "5", "#C56E6E"),
    ]
    targets = [
        (sb_x + 160, sb_y + 82),
        (sb_x + 160, sb_y + 145),
        (sb_x + 160, sb_y + 212),
        (sb_x + 160, sb_y + 272),
        (sb_x + 160, sb_y + 332),
    ]
    labels = [
        "Поле плана",
        "Папка датасета",
        "Загрузка",
        "Камера",
        "Снимок",
    ]
    for (box, num, color), target, label in zip(callouts, targets, labels):
        rounded_box(draw, box, color, color, radius=18, width=2)
        draw.text((box[0] + 24, box[1] + 14), num, font=label_font, fill="white")
        draw.text((box[0] + 56, box[1] + 17), label, font=body_font, fill="white")
        draw_arrow(draw, (box[0] + 10, (box[1] + box[3]) // 2), target, color, width=5)

    draw.rounded_rectangle((70, 875, 1730, 920), radius=16, fill="#F1F5F9", outline="#D7E3F1", width=2)
    draw.text((95, 885), "Зелёные элементы = данные для сохранения. Синие элементы = навигация и действия оператора.", font=body_font, fill="#475569")

    path.parent.mkdir(parents=True, exist_ok=True)
    img.save(path)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width) -> None:
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.twips)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "8")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "D7E3F1")
        borders.append(elem)
    tblPr.append(borders)


def set_paragraph_spacing(paragraph, before=0, after=0, line=1.25):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def style_run(run, name="Calibri", size=11, bold=False, color="000000", italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def apply_base_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color in [("Heading 1", 16, "2E74B5"), ("Heading 2", 13, "2E74B5"), ("Heading 3", 12, "1F4D78")]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)


def add_title(doc: Document, text: str, subtitle: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    style_run(r, size=26, bold=False, color="000000")

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_after = Pt(10)
    r2 = p2.add_run(subtitle)
    style_run(r2, size=11, color="555555")


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        p.add_run(item)


def add_note_box(doc: Document, text: str, label: str = "Важно") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    set_table_borders(table)
    cell = table.cell(0, 0)
    set_cell_width(cell, Inches(6.5))
    set_cell_shading(cell, "F4F6F9")
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run(f"{label}. ")
    style_run(r1, bold=True, color="1F4D78")
    r2 = p.add_run(text)
    style_run(r2, color="000000")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], col_widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    set_table_borders(table)
    for idx, width in enumerate(col_widths):
        table.columns[idx].width = Inches(width)

    hdr = table.rows[0].cells
    for i, head in enumerate(headers):
        cell = hdr[i]
        cell.text = ""
        set_cell_width(cell, Inches(col_widths[i]))
        set_cell_shading(cell, "E8EEF5")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(head)
        style_run(run, bold=True, color="17324D")

    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cell = cells[i]
            cell.text = ""
            set_cell_width(cell, Inches(col_widths[i]))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            style_run(run, color="000000")

    doc.add_paragraph()


def add_image(doc: Document, path: Path, width_inches: float, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Inches(width_inches))
    p.paragraph_format.space_after = Pt(4)

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    rr = cap.add_run(caption)
    style_run(rr, size=9, italic=True, color="555555")


def build_doc() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    make_workflow_image(WORKFLOW_IMG)
    make_ui_map_image(UI_MAP_IMG)

    doc = Document()
    apply_base_styles(doc)

    # Header / footer
    section = doc.sections[0]
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer_p.add_run("CassetteDatasetCapture • Инструкция оператора")
    style_run(footer_run, size=9, color="6B7280")

    add_title(
        doc,
        "Инструкция оператора CassetteDatasetCapture",
        "Короткая памятка для запуска камеры, загрузки плана и сохранения датасета",
    )

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Назначение: ")
    style_run(r, bold=True, color="1F4D78")
    r2 = p.add_run("оператор последовательно снимает кассеты по плану и сохраняет изображения в выбранную папку датасета.")
    style_run(r2)

    add_note_box(
        doc,
        "Для реальной камеры выберите режим Daheng. Для обучения и проверки интерфейса используйте режим «Тестовая».",
    )

    doc.add_heading("1. Быстрый старт", level=1)
    add_numbered(
        doc,
        [
            "Выберите режим камеры.",
            "Нажмите «Обзор» рядом с полем конфигурации и укажите JSON-файл плана.",
            "Нажмите «Обзор» рядом с папкой датасета и выберите каталог для сохранения снимков.",
            "Нажмите «Загрузить план».",
            "Нажмите «Подключить камеру».",
            "Запустите просмотр, если нужно видеть живой кадр.",
            "Нажмите «Сделать снимок» для сохранения текущего кадра.",
            "Если кадр нужно заменить, используйте «Переснять».",
            "Если текущую конфигурацию нужно пропустить, нажмите «Пропустить».",
        ],
    )

    add_image(
        doc,
        WORKFLOW_IMG,
        6.5,
        "Рисунок 1. Основной порядок действий оператора с направлением движения по шагам",
    )

    doc.add_heading("2. Что находится в окне программы", level=1)
    add_image(
        doc,
        UI_MAP_IMG,
        6.5,
        "Рисунок 2. Схема интерфейса и основные зоны, на которые нужно смотреть",
    )

    doc.add_heading("3. Что делает каждый элемент", level=1)
    add_table(
        doc,
        ["Элемент", "Назначение"],
        [
            ["Путь к конфигурации", "Путь к JSON-файлу плана съёмки. Можно выбрать через кнопку «Обзор»."],
            ["Папка датасета", "Каталог, куда программа будет складывать изображения и служебные файлы."],
            ["Загрузить план", "Читает конфигурацию и строит список задач для оператора."],
            ["Подключить камеру", "Открывает выбранный режим камеры и готовит её к работе."],
            ["Запустить просмотр", "Показывает живой видеопоток в окне предпросмотра."],
            ["Сделать снимок", "Сохраняет текущий кадр в датасет и продвигает счётчик снимков."],
            ["Переснять", "Помечает предыдущий кадр как rejected и делает новый снимок."],
            ["Пропустить", "Пропускает текущую конфигурацию, если она не подходит для съёмки."],
            ["Статус внизу окна", "Показывает ошибки, подсказки и текущее состояние работы."],
        ],
        [1.85, 4.65],
    )

    doc.add_heading("4. Перед началом съёмки", level=1)
    add_bullets(
        doc,
        [
            "Проверьте, что выбран правильный режим камеры.",
            "Убедитесь, что путь к плану указывает на нужный JSON-файл.",
            "Проверьте, что папка датасета доступна для записи.",
            "Посмотрите на подсказку в центре окна и на схему слотов.",
            "Если что-то не совпадает с заданием, сначала исправьте конфигурацию, а затем начинайте съёмку.",
        ],
    )

    doc.add_heading("5. Где сохраняются результаты", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("По умолчанию программа создаёт внутри выбранной папки датасета структуру вида: ")
    style_run(r)
    r2 = p.add_run("images / progress.json")
    style_run(r2, bold=True, color="1F4D78")
    add_note_box(
        doc,
        "Если папка датасета пустая, программа создаст нужные подпапки автоматически после загрузки плана.",
        label="Подсказка",
    )

    doc.add_heading("6. Если что-то пошло не так", level=1)
    add_bullets(
        doc,
        [
            "Если камера не подключается, проверьте кабель, питание и выбранный режим камеры.",
            "Если план не загружается, проверьте путь к JSON-файлу и его содержимое.",
            "Если снимки не сохраняются, проверьте права на запись в папку датасета.",
            "Если изображение в предпросмотре не меняется, остановите просмотр и запустите его снова.",
        ],
    )

    out = OUT_DOCX
    doc.save(out)


if __name__ == "__main__":
    build_doc()
